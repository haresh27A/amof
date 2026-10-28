"""Generate an editable PowerPoint version of the AMORA architecture.

Every box, label and arrow is a native PowerPoint shape. Arrows between boxes are
glued connectors, so they follow when boxes are moved.
Export to PDF for LaTeX: File > Export > PDF (or Save As > PDF).
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "amora_architecture.pptx")
U = 0.8          # inches per layout unit
TOP = 14.0
W_UNITS, H_UNITS = 16.0, 14.4

C = {
    "krt": ("FBE9E7", "D84315"),
    "kqf": ("E3F2FD", "1565C0"),
    "route": ("EDE7F6", "4527A0"),
    "fuse": ("E8F5E9", "2E7D32"),
    "gen": ("FFF8E1", "EF6C00"),
    "neutral": ("FFFFFF", "546E7A"),
    "query": ("FFFFFF", "7CB342"),
}
INK = RGBColor(0x26, 0x32, 0x38)
GREY = RGBColor(0x78, 0x90, 0x9C)
FONT = "Calibri"

prs = Presentation()
prs.slide_width = Inches(W_UNITS * U)
prs.slide_height = Inches(H_UNITS * U)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


def P(x, y):
    """layout units (y up) -> EMU (y down)"""
    return Inches(x * U), Inches((TOP - y) * U)


def rgb(h):
    return RGBColor.from_string(h)


def set_text(tf, title, lines, title_color, size=11, align=PP_ALIGN.CENTER):
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    first = True
    items = ([(title, True)] if title else []) + [(l, False) for l in lines]
    for text, bold in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size + 1 if bold else size)
        r.font.bold = bold
        r.font.color.rgb = rgb(title_color) if bold else INK


def box(cx, cy, w, h, title, lines, kind, dashed=False, align=PP_ALIGN.CENTER, size=11):
    fc, ec = C[kind]
    x, y = P(cx - w / 2, cy + h / 2)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(w * U), Inches(h * U))
    s.adjustments[0] = 0.12
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fc)
    s.line.color.rgb = rgb(ec)
    s.line.width = Pt(1.75)
    if dashed:
        s.line.dash_style = MSO_LINE.DASH
    s.shadow.inherit = False
    set_text(s.text_frame, title, lines, ec, size, align)
    return s


def tag(x, y, text, kind, w=None):
    fc, ec = C[kind]
    w = w or 0.105 * len(text) + 0.35
    gx, gy = P(x - 0.15, y + 0.2)
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, gx, gy, Inches(w * U), Inches(0.4 * U))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fc)
    s.line.color.rgb = rgb(ec)
    s.line.width = Pt(1.25)
    s.shadow.inherit = False
    set_text(s.text_frame, text, [], ec, 13)
    s.text_frame.word_wrap = False


def label(cx, cy, lines, color=INK, size=10, w=2.6, bold=True, rot=0):
    x, y = P(cx - w / 2, cy + 0.3)
    tb = slide.shapes.add_textbox(x, y, Inches(w * U), Inches(0.6 * U))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, t in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = t
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    tb.rotation = rot
    return tb


def _style_line(conn, arrow=True, dashed=False, color="263238", width=1.75, dash=MSO_LINE.DASH):
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width)
    if dashed:
        conn.line.dash_style = dash
    if arrow:
        ln = conn.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")


# connection-site indices for (rounded) rectangles: 0 top, 1 left, 2 bottom, 3 right
TOPC, LEFT, BOTTOM, RIGHT = 0, 1, 2, 3


def glue(a, ai, b, bi, kind=MSO_CONNECTOR.STRAIGHT, **kw):
    c = slide.shapes.add_connector(kind, 0, 0, 0, 0)
    c.begin_connect(a, ai)
    c.end_connect(b, bi)
    _style_line(c, **kw)
    return c


def seg(p, q, **kw):
    (x1, y1), (x2, y2) = P(*p), P(*q)
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    _style_line(c, **kw)
    return c


def sep(y):
    seg((0.2, y), (15.8, y), arrow=False, dashed=True, color="9E9E9E", width=1.5)


# ---------------------------------------------------------------- Query ----
query = box(8.3, 13.43, 14.6, 0.62, None,
            ["Query:  What is the current stock of SKU-1042 in the Chennai warehouse, "
             "and does it match the Q3 inventory report?"], "query", dashed=True,
            align=PP_ALIGN.LEFT, size=12)
sep(12.8)

# ---------------------------------------------------------------- IE-KRT ---
tag(2.0, 12.45, "IE-KRT", "krt")
ev = box(3.75, 11.7, 2.9, 1.05, "Evaluator", ["Complexity  E = 0.72", "Sensitivity  Sᵢ = 0.76"], "krt")
sa = box(3.75, 10.55, 2.9, 0.85, "LLM Self-Assessment", ["Confidence  mᵢ = 0.18"], "krt")
ks = box(8.0, 11.1, 3.5, 1.0, "Activation Score", ["K(tᵢ) = (E − mᵢ) · Sᵢ = 0.41"], "neutral")
pr = box(13.4, 11.1, 3.4, 1.0, "Parametric Response",
         ["Direct LLM answer, 0 ms retrieval", "(35.0% of queries bypassed)"], "neutral", dashed=True)
glue(ev, RIGHT, ks, LEFT)
glue(sa, RIGHT, ks, LEFT)
glue(ks, RIGHT, pr, LEFT)
label(10.72, 11.42, ["K < θ"], size=10)
label(10.72, 10.78, ["No Need Retrieval"], size=9)
sep(9.9)

# ---------------------------------------------------------------- IE-KQF ---
tag(2.0, 9.6, "IE-KQF", "kqf")
ph = box(3.95, 8.4, 3.3, 1.3, "Internal Phrases (LLM)",
         ["SKU-1042, Chennai warehouse,", "Q3 inventory report"], "kqf")
it = box(8.0, 8.4, 3.3, 1.3, "Intent Tokens", ["live stock lookup +", "cross-source verification"], "kqf")
qp = box(13.1, 8.4, 4.0, 1.3, "Reformulated Query q⁺",
         ["current stock SKU-1042 Chennai", "warehouse | compare Q3 inventory report"], "kqf")
glue(ks, BOTTOM, it, TOPC)
label(9.4, 9.45, ["K ≥ θ = 0.25", "Need Retrieval"], size=9.5, w=2.2)
glue(ph, RIGHT, it, LEFT)
glue(it, RIGHT, qp, LEFT)
label(10.38, 8.7, ["concat"], color=GREY, size=9)

# Query trunk (vertical line + branches)
XT = 1.7
seg((XT, 13.12), (XT, 8.4), arrow=False)
seg((XT, 11.7), (2.3, 11.7))
seg((XT, 10.55), (2.3, 10.55))
seg((XT, 8.4), (2.3, 8.4))
sep(7.1)

# --------------------------------------------------------- AMORA Routing ---
tag(2.0, 6.8, "AMORA Routing", "route", w=2.15)
fn = box(13.1, 5.6, 4.0, 1.3, "Two-Stage FAISS Funnel",
         ["N = 10,000 → M = 55 candidates", "3.26 ms  |  SRR = 99.45%"], "route")
mo = box(8.0, 5.6, 3.3, 1.3, "6-D Multi-Objective Scoring",
         ["R, L, F, C, H, K", "Σⱼ Wⱼ = 1 (Pareto-optimal)"], "route")
tk = box(3.95, 5.6, 3.3, 1.3, "Top-K Selected Tools",
         ["inventory_db (SQL)", "stock_api (REST), q3_report (DOC)"], "route")
glue(qp, BOTTOM, fn, TOPC)
label(13.85, 7.0, ["Retrieval"], color=GREY, size=9, w=1.4)
glue(fn, LEFT, mo, RIGHT)
glue(mo, LEFT, tk, RIGHT)
sep(4.4)

# ------------------------------------------------------- Evidence Fusion ---
tag(2.0, 4.1, "Evidence Fusion", "fuse")
ou = box(3.95, 2.9, 3.3, 1.3, "Tool Outputs",
         ["SQL: 150 units (2 h ago)", "REST: 150 units (live)", "DOC: 90 units (Q3 snapshot)"],
         "fuse", size=10)
en = box(8.0, 2.9, 3.3, 1.3, "Fusion Engine",
         ["Hash-window dedup: 3 → 2 facts", "O(E²) conflict graph: 150 vs 90",
          "Prompt tokens −26.7%"], "fuse", size=10)
fu = box(13.1, 2.9, 4.0, 1.3, "Fused Evidence Context",
         ["[FACT] stock = 150 (fresh, reliable)", "[CONFLICT] Q3 report = 90 (stale)",
          "provenance tags attached"], "fuse", size=10)
glue(tk, BOTTOM, ou, TOPC)
glue(ou, RIGHT, en, LEFT)
glue(en, RIGHT, fu, LEFT)
sep(1.85)

# ------------------------------------------------------------ Generation ---
lm = box(13.1, 0.6, 3.0, 1.05, "Local LLM", ["Qwen 2.5 7B (Q4_K_M)", "zero data egress"], "gen")
an = box(6.25, 0.6, 7.7, 1.05, None,
         ["SKU-1042 currently has 150 units in the Chennai warehouse [inventory_db, updated 2 h ago].",
          "Note: the Q3 report lists 90 units; the gap reflects post-Q3 restocking,",
          "so the live ledger value is authoritative."], "query", align=PP_ALIGN.LEFT, size=10)
glue(fu, BOTTOM, lm, TOPC)
glue(lm, LEFT, an, RIGHT)
label(6.25, -0.15, ["LLM Continual Generation based on Fused Evidence"], color=GREY, size=10, w=7.0)

# Iteration loop (dashed)
seg((2.4, 0.6), (0.7, 0.6), arrow=False, dashed=True)
seg((0.7, 0.6), (0.7, 12.2), arrow=False, dashed=True)
seg((0.7, 12.2), (XT, 12.2), dashed=True)
label(0.35, 6.3, ["↻ Iteration"], color=RGBColor(0x15, 0x65, 0xC0), size=12, w=2.0, rot=270)

prs.save(OUT)
print("Saved", OUT)
